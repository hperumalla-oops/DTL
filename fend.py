import streamlit as st
import requests
from PyPDF2 import PdfReader
from pptx import Presentation
from docx import Document
from PIL import Image

# Backend URLs
backend_upload_url = "http://127.0.0.1:5000/upload"
backend_verify_url = "http://127.0.0.1:5000/verify"

def count_pages(file, file_extension):
    try:
        if file_extension == "pdf":
            pdf_reader = PdfReader(file)
            return len(pdf_reader.pages)
        elif file_extension == "pptx":
            presentation = Presentation(file)
            return len(presentation.slides)
        elif file_extension == "docx":
            document = Document(file)
            return len(document.paragraphs)  # Approximation for Word files
        elif file_extension == "png":
            img = Image.open(file)
            img.verify()  # Verify if the image is valid
            return 1
        else:
            return -1  # Unsupported file
    except Exception as e:
        st.error(f"Error processing file: {e}")
        return -1

def calculate_cost(specs):
    page_count = specs['page_count']
    size = specs['size']
    color = specs['color']
    side = specs['side']
    binding = specs['binding']
    paper = specs['paper']
    lamination = specs['lamination']
    copies = specs['copies']

    cost = 0
    if size == 'A3':
        cost = 40 * page_count
    elif size == "Passport":
        cost = 8
    else:
        if color == 'Colored':
            cost = page_count * 5
        else:
            if side == 'Single':
                cost = page_count
            else:
                cost = page_count * 0.75
    if binding == 'Tape':
        cost += 5
    elif binding == 'Spiral':
        cost += 20
    if paper == "Bond":
        cost += page_count * 3
    if lamination == 'Yes':
        cost += 30

    return cost * copies

st.title("File Upload with Specifications, Cost Calculation, and Verification")

uploaded_files = st.file_uploader(
    "Upload your files (PDF, PNG, PPTX, DOCX)", 
    type=["pdf", "png", "pptx", "docx"], 
    accept_multiple_files=True
)

if uploaded_files:
    #file_extension = uploaded_files.name.split('.')[-1].lower()
    file_details = []
    total_cost = 0

    for idx, uploaded_file in enumerate(uploaded_files):
        file_extension = uploaded_file.name.split('.')[-1].lower()
        page_count = count_pages(uploaded_file, file_extension)

        st.write("Page count:",page_count)
        st.write(f"**Specifications for {uploaded_file.name}:**")
        binding = st.selectbox(f"Binding for {uploaded_file.name}", ["None", "Spiral", "Tape"], key=f"binding_{idx}")
        paper = st.selectbox(f"Paper for {uploaded_file.name}", ["Regular", "Bond"], key=f"paper_{idx}")
        size = st.selectbox(f"Size for {uploaded_file.name}", ["A4", "A3", "Passport"], key=f"size_{idx}")
        color = st.selectbox(f"Color for {uploaded_file.name}", ["Colored", "Black-and-White"], key=f"color_{idx}")
        side = st.selectbox(f"Side for {uploaded_file.name}", ["Double", "Single"], key=f"side_{idx}")
        lamination = st.selectbox(f"Lamination for {uploaded_file.name}", ["No", "Yes"], key=f"lamination_{idx}")
        copies = st.number_input(f"Copies for {uploaded_file.name}", min_value=1, step=1, key=f"copies_{idx}")

        specs = {
            "name": uploaded_file.name,
            "type": file_extension.upper(),
            "page_count": page_count if page_count > 0 else "Unable to compute",
            "binding": binding,
            "paper": paper,
            "size": size,
            "color": color,
            "side": side,
            "lamination": lamination,
            "copies": copies
        }
        file_cost = calculate_cost(specs)
        #specs['cost'] = file_cost

        total_cost += file_cost
        file_details.append(specs)

        st.write(f"**Cost for {uploaded_file.name}: ₹{file_cost}**")
        st.write("---")

    for file_spec in file_details:
        file_spec['cost'] = total_cost

    st.write(f"### Total Cost for All Files: ₹{total_cost}")

    if st.button("Upload All Files"):
        with st.spinner("Uploading..."):
            files_to_send = [('files', (file.name, file, file.type)) for file in uploaded_files]
            metadata = str(file_details)
            response = requests.post(
                backend_upload_url,
                files=files_to_send,
                data={'metadata': metadata})

        if response.status_code == 200:
            st.success(f"Files uploaded successfully!"
                       f"The total cost of your printouts is ₹{total_cost}. "
                        f"Please pay ₹{total_cost} to the following phone number: 99XXXXXXXXX (MR. Manoj, RVCE Printing). "
                        f"After payment, click on verify."
                )
        else:
            st.write(f"Error uploading files: {response.json()}")

    if st.button("Verify"):
        with st.spinner("Verifying..."):
            response = requests.post(backend_verify_url)
            if response.status_code == 200:
                st.success("Verification successful!")
            else:
                st.error("Verification failed!")


